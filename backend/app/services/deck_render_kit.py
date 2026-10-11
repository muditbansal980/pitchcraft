"""Drawing primitives for the PPTX renderer (Agent B).

python-pptx exposes solid fills, lines and text; transparency, multi-stop
gradients, letter spacing and table borders need a little DrawingML XML, which
lives here so the layout builders in ``deck_render`` stay readable.

Text fitting uses each theme font's real advance widths
(``assets/font_metrics.json``, measured from the Google Fonts files with Pillow;
widths only, no font data) with ~5% headroom for substitution, a greedy word wrap
like PowerPoint's, and font sizes that step down until the wrapped text fits.
"""

import json
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

_METRICS_PATH = Path(__file__).resolve().parent.parent / "assets" / "font_metrics.json"
# Headroom for PowerPoint/Canva substituting a font or rendering slightly wider.
HEADROOM = 1.05
_UNKNOWN_EM = 0.62  # emoji, Devanagari, anything outside the tables


def _load_metrics() -> tuple[dict[str, dict[str, dict[str, float]]], dict[str, float]]:
    try:
        raw = json.loads(_METRICS_PATH.read_text())
        chars = raw["chars"]
        fonts = {
            fam: {style: {c: w / 1000 for c, w in zip(chars, widths, strict=True)} for style, widths in styles.items()}
            for fam, styles in raw["fonts"].items()
        }
    except (OSError, ValueError, KeyError):
        fonts = {}
    fallback = {chr(c): 0.56 for c in range(32, 127)}
    fallback[" "] = 0.27
    return fonts, fallback


_FONTS, _FALLBACK = _load_metrics()


_CATEGORY_STAND_IN = {
    "geometric": "montserrat", "humanist": "open sans", "grotesk": "inter", "condensed": "oswald",
    "slab": "roboto slab", "serif": "lora", "didone": "playfair display", "rounded": "nunito",
    "script": "pacifico", "display": "archivo black",
}


def _widths(font: str, bold: bool) -> dict[str, float]:
    key = (font or "").strip().lower()
    fam = _FONTS.get(key)
    if not fam:  # an unknown font: use a stand-in of the same kind (condensed, serif, ...)
        from app.services.brand_identity_fonts import classify

        fam = _FONTS.get(_CATEGORY_STAND_IN.get(classify(key) or "", "montserrat")) or _FONTS.get("montserrat")
    if not fam:
        return _FALLBACK
    return fam.get("bold" if bold else "regular") or fam.get("regular") or _FALLBACK


ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


# --- text measuring ---------------------------------------------------------


def measure(text: str, size_pt: float, font: str, *, bold: bool = False, upper: bool = False) -> float:
    """Width of one line of ``text`` in inches."""
    table = _widths(font, bold)
    if upper:
        text = text.upper()
    em = sum(table.get(ch, _UNKNOWN_EM) for ch in text)
    return em * size_pt / 72.0 * HEADROOM


def text_width(text: str, size_pt: float, font: str, upper: bool = False, bold: bool = False) -> float:
    return measure(text, size_pt, font, bold=bold, upper=upper)


def wrap_lines(text: str, width_in: float, size_pt: float, font: str, *, bold: bool = False,
               upper: bool = False) -> list[str]:
    """Greedy word wrap using the font's advance widths (``\n`` starts a paragraph)."""
    table = _widths(font, bold)
    scale = size_pt / 72.0 * HEADROOM
    space = table.get(" ", 0.27) * scale

    def w(s: str) -> float:
        return sum(table.get(ch, _UNKNOWN_EM) for ch in (s.upper() if upper else s)) * scale

    lines: list[str] = []
    for para in (text or "").split("\n"):
        words = para.split()
        if not words:
            lines.append("")
            continue
        cur, cur_w = "", 0.0
        for word in words:
            ww = w(word)
            if not cur:
                cur, cur_w = word, ww
            elif cur_w + space + ww <= width_in:
                cur, cur_w = f"{cur} {word}", cur_w + space + ww
            else:
                lines.append(cur)
                cur, cur_w = word, ww
            while cur_w > width_in and len(cur) > 1:  # a word longer than the line breaks
                cut = max(1, int(len(cur) * width_in / cur_w))
                lines.append(cur[:cut])
                cur = cur[cut:]
                cur_w = w(cur)
        lines.append(cur)
    return lines


def wrap_count(text: str, width_in: float, size_pt: float, font: str, *, bold: bool = False,
               upper: bool = False) -> int:
    if not text:
        return 0
    return len(wrap_lines(text, width_in, size_pt, font, bold=bold, upper=upper))


def clip_lines(text: str, width_in: float, size_pt: float, font: str, max_lines: int, *, bold: bool = False,
               upper: bool = False) -> str:
    """``text`` cut (with an ellipsis) so it wraps to at most ``max_lines`` lines."""
    if max_lines <= 0:
        return ""
    lines = wrap_lines(text, width_in, size_pt, font, bold=bold, upper=upper)
    if len(lines) <= max_lines:
        return text
    kept = lines[:max_lines]
    last = kept[-1]
    while last and measure(last + "…", size_pt, font, bold=bold, upper=upper) > width_in:
        last = last.rsplit(" ", 1)[0] if " " in last else last[:-1]
    kept[-1] = last.rstrip(" ,;:.") + "…"
    return "\n".join(kept) if "\n" in text else " ".join(kept)


def clip_to_box(text: str, width_in: float, height_in: float, size_pt: float, font: str, *, spacing: float,
                bold: bool = False, upper: bool = False) -> str:
    return clip_lines(text, width_in, size_pt, font, max(1, int(height_in / (size_pt * spacing / 72.0) + 1e-6)),
                      bold=bold, upper=upper)


def fit_or_clip(text: str, width_in: float, height_in: float, font: str, *, max_pt: float, min_pt: float,
                spacing: float = 1.2, bold: bool = False, upper: bool = False) -> tuple[str, "Fit"]:
    """Like ``fit_text``; if the text still overflows at ``min_pt`` it is clipped to the box."""
    f = fit_text(text, width_in, height_in, font, max_pt=max_pt, min_pt=min_pt, spacing=spacing, bold=bold,
                 upper=upper)
    if f.height <= height_in + 1e-6:
        return text, f
    clipped = clip_to_box(text, width_in, height_in, f.size, font, spacing=spacing, bold=bold, upper=upper)
    lines = wrap_count(clipped, width_in, f.size, font, bold=bold, upper=upper)
    return clipped, Fit(size=f.size, lines=lines, line_h=f.line_h)


@dataclass(frozen=True)
class Fit:
    size: float
    lines: int
    line_h: float  # inches per line

    @property
    def height(self) -> float:
        return self.lines * self.line_h


def fit_text(
    text: str,
    width_in: float,
    height_in: float,
    font: str,
    *,
    max_pt: float,
    min_pt: float,
    upper: bool = False,
    bold: bool = False,
    spacing: float = 1.2,
    max_lines: int | None = None,
) -> Fit:
    """Largest size in [min_pt, max_pt] at which ``text`` fits; ``min_pt`` if nothing fits."""
    size = max_pt
    while True:
        lines = wrap_count(text, width_in, size, font, bold=bold, upper=upper)
        lh = size * spacing / 72.0
        fits = lines * lh <= height_in and (max_lines is None or lines <= max_lines)
        if fits or size <= min_pt:
            return Fit(size=max(size, min_pt), lines=max(lines, 1), line_h=lh)
        size = max(min_pt, size - (0.5 if size < 16 else 1.0 if size < 32 else 2.0))


# --- XML helpers ------------------------------------------------------------


def _srgb(hex_color: str, alpha: float | None = None) -> etree._Element:
    el = etree.Element(qn("a:srgbClr"), val=hex_color.lstrip("#").upper())
    if alpha is not None and alpha < 1:
        etree.SubElement(el, qn("a:alpha"), val=str(int(max(0, min(1, alpha)) * 100000)))
    return el


def _replace_fill(sp_pr: etree._Element, fill_el: etree._Element) -> None:
    for tag in ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
        for old in sp_pr.findall(qn(tag)):
            sp_pr.remove(old)
    # Fill goes right after the geometry element.
    geom = sp_pr.find(qn("a:prstGeom"))
    if geom is None:
        geom = sp_pr.find(qn("a:custGeom"))
    if geom is not None:
        geom.addnext(fill_el)
    else:
        sp_pr.insert(0, fill_el)


def set_solid(shape, hex_color: str, alpha: float | None = None) -> None:
    fill = etree.Element(qn("a:solidFill"))
    fill.append(_srgb(hex_color, alpha))
    _replace_fill(shape._element.spPr, fill)


def set_gradient(shape, stops: Sequence[tuple[float, str, float]], angle: float = 0.0) -> None:
    """Linear gradient; ``stops`` = (position 0..1, hex, alpha 0..1); ``angle`` in degrees (0 = left->right)."""
    grad = etree.Element(qn("a:gradFill"), rotWithShape="1")
    gs_lst = etree.SubElement(grad, qn("a:gsLst"))
    for pos, color, alpha in stops:
        gs = etree.SubElement(gs_lst, qn("a:gs"), pos=str(int(max(0, min(1, pos)) * 100000)))
        gs.append(_srgb(color, alpha))
    etree.SubElement(grad, qn("a:lin"), ang=str(int(angle * 60000) % 21600000), scaled="0")
    _replace_fill(shape._element.spPr, grad)


def no_line(shape) -> None:
    shape.line.fill.background()


def set_line(shape, hex_color: str | None, width_pt: float = 1.0, dash: bool = False) -> None:
    if not hex_color:
        no_line(shape)
        return
    shape.line.color.rgb = rgb(hex_color)
    shape.line.width = Pt(width_pt)
    if dash:
        shape.line.dash_style = MSO_LINE_DASH_STYLE.DASH


def no_shadow(shape) -> None:
    sp_pr = shape._element.spPr
    if sp_pr.find(qn("a:effectLst")) is None:
        etree.SubElement(sp_pr, qn("a:effectLst"))


# --- shapes -----------------------------------------------------------------


def box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: str | None,
    *,
    alpha: float | None = None,
    line: str | None = None,
    line_w: float = 1.0,
    dash: bool = False,
    shape: MSO_SHAPE = MSO_SHAPE.RECTANGLE,
    radius: float | None = None,
    adj: float | None = None,
):
    """A filled shape; ``radius`` (inches) for rounded rectangles, ``adj`` for other presets."""
    if radius is not None:
        shape = MSO_SHAPE.ROUNDED_RECTANGLE
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(max(w, 0.01)), Inches(max(h, 0.01)))
    if radius is not None and min(w, h) > 0:
        shp.adjustments[0] = max(0.0, min(0.5, radius / min(w, h)))
    elif adj is not None and len(shp.adjustments):
        shp.adjustments[0] = adj
    if fill:
        set_solid(shp, fill, alpha)
    else:
        shp.fill.background()
    set_line(shp, line, line_w, dash)
    no_shadow(shp)
    if shp.has_text_frame:
        shp.text_frame.text = ""
    return shp


# --- text -------------------------------------------------------------------


@dataclass
class Run:
    text: str
    color: str | None = None
    bold: bool | None = None
    italic: bool | None = None
    font: str | None = None
    size: float | None = None


def _letter_spacing(run, spc_pt: float) -> None:
    run._r.get_or_add_rPr().set("spc", str(int(spc_pt * 100)))


def text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    content: str | Sequence[Sequence[Run]],
    *,
    font: str,
    size: float,
    color: str,
    bold: bool = False,
    italic: bool = False,
    align: str = "left",
    anchor: str = "top",
    spacing: float = 1.2,
    upper: bool = False,
    tracking: float = 0.0,
    space_after: float = 0.0,
    wrap: bool = True,
    autofit: bool = False,
):
    """A text box with zero insets. ``content`` is a string (``\\n`` = new paragraph) or paragraphs of Runs.

    Line spacing is set as an exact point value (``size * spacing``) so fonts with
    unusual metrics (Anton, Bebas) lay out the same in PowerPoint, LibreOffice and Canva.
    """
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(max(w, 0.1)), Inches(max(h, 0.1)))
    tf = tb.text_frame
    tf.word_wrap = wrap
    # ``autofit``: shrink-on-overflow as a safety net when a viewer substitutes a wider font.
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE if autofit else MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = ANCHOR[anchor]

    if isinstance(content, str):
        paragraphs = [[Run(line)] for line in content.split("\n")]
    else:
        paragraphs = [list(p) for p in content]

    for i, runs in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN[align]
        para.line_spacing = Pt(size * spacing)
        if space_after:
            para.space_after = Pt(space_after)
        for r in runs:
            if not r.text:
                continue
            run = para.add_run()
            run.text = r.text.upper() if upper else r.text
            f = run.font
            f.name = r.font or font
            f.size = Pt(r.size or size)
            f.bold = bold if r.bold is None else r.bold
            f.italic = italic if r.italic is None else r.italic
            f.color.rgb = rgb(r.color or color)
            if tracking and r.text.isascii():  # spacing breaks Devanagari/emoji shaping
                _letter_spacing(run, tracking)
    return tb


# --- pictures ---------------------------------------------------------------


def image_size(path: Path) -> tuple[int, int] | None:
    try:
        with Image.open(path) as img:
            return img.size
    except Exception:
        return None


def picture(slide, path: Path, x: float, y: float, w: float, h: float, *, mode: str = "cover",
            top_bias: float = 0.35):
    """Place an image in a box. ``cover`` crops to fill (biased toward the top, where faces
    and meme captions usually are); ``contain`` letterboxes and returns the fitted box."""
    size = image_size(path)
    if not size:
        return None, (x, y, w, h)
    iw, ih = size
    ia, ba = iw / ih, w / h
    if mode == "contain":
        if ia > ba:
            nw, nh = w, w / ia
        else:
            nw, nh = h * ia, h
        nx, ny = x + (w - nw) / 2, y + (h - nh) / 2
        pic = slide.shapes.add_picture(str(path), Inches(nx), Inches(ny), Inches(nw), Inches(nh))
        return pic, (nx, ny, nw, nh)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    if ia > ba:
        c = (1 - ba / ia) / 2
        pic.crop_left = pic.crop_right = c
    elif ia < ba:
        c = 1 - ia / ba
        pic.crop_top = c * top_bias
        pic.crop_bottom = c * (1 - top_bias)
    return pic, (x, y, w, h)


def set_notes(slide, notes: str) -> None:
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


# --- tables -----------------------------------------------------------------

NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def table_cell_style(cell, *, fill: str, border: str, border_w: float = 0.75, margin: float = 0.08,
                     margin_y: float | None = None) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = rgb(fill)
    cell.margin_left = cell.margin_right = Inches(margin)
    cell.margin_top = cell.margin_bottom = Inches(margin * 0.7 if margin_y is None else margin_y)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tc_pr.findall(qn(tag)):
            tc_pr.remove(old)
        ln = etree.Element(qn(tag), w=str(int(border_w * 12700)), cap="flat", cmpd="sng", algn="ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        solid.append(_srgb(border))
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        # Borders must precede the cell fill in tcPr.
        tc_pr.insert(len([c for c in tc_pr if c.tag in {qn(t) for t in ("a:lnL", "a:lnR", "a:lnT", "a:lnB")}]), ln)


def set_table_style(table, style_id: str = NO_STYLE_TABLE) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_pr.set("firstRow", "1")
    tbl_pr.set("bandRow", "0")
    sid = tbl_pr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tbl_pr, qn("a:tableStyleId"))
    sid.text = style_id
